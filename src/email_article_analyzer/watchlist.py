from dataclasses import dataclass
import csv
from io import BytesIO, StringIO
import json
import re
from xml.etree import ElementTree
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook


@dataclass(frozen=True)
class ParsedWatchlist:
    columns: list[str]
    rows: list[dict[str, str]]


def normalize_ticker(value: str) -> str:
    ticker = value.strip().upper()
    ticker = ticker.removeprefix("$")
    if ":" in ticker:
        ticker = ticker.split(":", 1)[1]
    return re.sub(r"[^A-Z0-9.\-]", "", ticker)


def parse_watchlist_file(filename: str, content: bytes) -> ParsedWatchlist:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        return _parse_csv(content)
    if lower_name.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise ValueError("Unsupported watchlist file type. Use CSV or XLSX.")


def _parse_csv(content: bytes) -> ParsedWatchlist:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    columns = list(reader.fieldnames or [])
    rows = [
        {column: (row.get(column) or "").strip() for column in columns}
        for row in reader
    ]
    return ParsedWatchlist(columns=columns, rows=rows)


def _parse_xlsx(content: bytes) -> ParsedWatchlist:
    workbook = load_workbook(
        BytesIO(_strip_conditional_formatting(content)),
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return ParsedWatchlist(columns=[], rows=[])
    columns = [str(cell or "").strip() for cell in values[0]]
    rows: list[dict[str, str]] = []
    for row in values[1:]:
        mapped = {
            columns[index]: str(row[index] or "").strip()
            for index in range(len(columns))
        }
        if any(mapped.values()):
            rows.append(mapped)
    return ParsedWatchlist(columns=columns, rows=rows)


def _strip_conditional_formatting(content: bytes) -> bytes:
    source = BytesIO(content)
    target = BytesIO()
    with ZipFile(source, "r") as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = _strip_conditional_formatting_xml(data)
            dst.writestr(item, data)
    return target.getvalue()


def _strip_conditional_formatting_xml(data: bytes) -> bytes:
    root = ElementTree.fromstring(data)
    namespace = ""
    if root.tag.startswith("{"):
        namespace = root.tag.split("}", 1)[0][1:]
    conditional_tag = f"{{{namespace}}}conditionalFormatting" if namespace else "conditionalFormatting"
    for child in list(root):
        if child.tag == conditional_tag:
            root.remove(child)
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


@dataclass(frozen=True)
class ValidWatchlistRow:
    raw_ticker: str
    normalized_ticker: str
    values: dict[str, str]
    polygon_reference: dict


@dataclass(frozen=True)
class InvalidWatchlistRow:
    raw_ticker: str
    values: dict[str, str]
    reason: str


@dataclass(frozen=True)
class WatchlistValidationResult:
    valid_rows: list[ValidWatchlistRow]
    invalid_rows: list[InvalidWatchlistRow]


class WatchlistService:
    def __init__(self, repository, polygon_validator):
        self.repository = repository
        self.polygon_validator = polygon_validator

    def validate(
        self,
        parsed: ParsedWatchlist,
        ticker_column: str,
        optional_columns: dict[str, str],
    ) -> WatchlistValidationResult:
        valid_rows: list[ValidWatchlistRow] = []
        invalid_rows: list[InvalidWatchlistRow] = []

        for row in parsed.rows:
            raw_ticker = row.get(ticker_column, "")
            normalized = normalize_ticker(raw_ticker)
            if not normalized:
                invalid_rows.append(InvalidWatchlistRow(raw_ticker, row, "Missing ticker"))
                continue
            polygon_reference = self.polygon_validator.validate_ticker(normalized)
            if polygon_reference is None:
                invalid_rows.append(
                    InvalidWatchlistRow(
                        raw_ticker,
                        row,
                        "Polygon did not recognize ticker",
                    )
                )
                continue
            valid_rows.append(
                ValidWatchlistRow(
                    raw_ticker=raw_ticker,
                    normalized_ticker=normalized,
                    values=row,
                    polygon_reference=polygon_reference,
                )
            )
        return WatchlistValidationResult(valid_rows=valid_rows, invalid_rows=invalid_rows)

    def replace_active_watchlist(
        self,
        original_filename: str,
        parsed: ParsedWatchlist,
        ticker_column: str,
        optional_columns: dict[str, str],
    ) -> WatchlistValidationResult:
        result = self.validate(parsed, ticker_column, optional_columns)
        upload_id = self.repository.create_upload(
            original_filename=original_filename,
            columns=parsed.columns,
            sample_rows=parsed.rows[:5],
            row_count=len(parsed.rows),
        )
        items = []
        for row in result.valid_rows:
            values = row.values
            items.append(
                {
                    "ticker": row.raw_ticker,
                    "normalized_ticker": row.normalized_ticker,
                    "company_name": _optional_value(values, optional_columns, "company_name"),
                    "sector": _optional_value(values, optional_columns, "sector"),
                    "priority": _optional_value(values, optional_columns, "priority"),
                    "notes": _optional_value(values, optional_columns, "notes"),
                    "polygon_reference": json.dumps(row.polygon_reference, sort_keys=True),
                }
            )
        self.repository.replace_items(upload_id, items)
        return result


def _optional_value(
    values: dict[str, str],
    optional_columns: dict[str, str],
    field_name: str,
) -> str | None:
    column_name = optional_columns.get(field_name)
    if not column_name:
        return None
    value = values.get(column_name)
    return value or None
